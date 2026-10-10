#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Batch CSV Converter
===================
Desktop app for converting many files to CSV in bulk.

Supported inputs:
    - Excel (.xlsx, .xlsm, .xls)
    - CSV (.csv)  -> standardised (delimiter detection and clean rewrite)
    - TXT / LOG   -> automatic delimiter detection and conversion to CSV
    - SQLite databases (.db, .sqlite, .sqlite3, .db3, or any file with a SQLite header)
    - Access databases (.mdb, .accdb) -> needs pyodbc + the Access driver (Windows)
    - .mdf files (SQL Server) -> help message only (a SQL Server engine is required)
    - SQL dumps (.sql) -> INSERT statements extracted as tables
    - .torrent files -> metadata extraction (name, trackers, file list)
    - PDF (.pdf) -> text extraction per page (needs pypdf)
    - Word (.docx) -> paragraphs and tables extracted (needs python-docx);
        legacy .doc files must be converted to .docx first
    - Any other file with any extension (or none) -> read as text, delimiter
        auto-detected and converted to CSV. If the file is binary/unreadable, only
        that file is flagged as failed and the rest are still processed.

Dependencies (install with pip):
    pip install openpyxl          # modern Excel
    pip install xlrd              # legacy .xls Excel (optional)
    pip install chardet           # better encoding detection (optional)
    pip install pyodbc            # Access .mdb/.accdb files (optional, Windows only)
    pip install pypdf             # PDF files (optional)
    pip install python-docx       # Word .docx files (optional)

Run:
    python batch_csv_converter.py

GitHub: https://github.com/Mrmmd2004/BatchCSVConverter
"""

import os
import sys
import csv
import io
import re
import queue
import sqlite3
import threading
import traceback
import webbrowser
import multiprocessing
from concurrent.futures import ProcessPoolExecutor, as_completed, wait as cf_wait, FIRST_COMPLETED as cf_FIRST_COMPLETED

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

# ---------------------------------------------------------------------------
# Optional dependencies
# ---------------------------------------------------------------------------
try:
    import chardet
    HAS_CHARDET = True
except ImportError:
    HAS_CHARDET = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

try:
    import pyodbc
    HAS_PYODBC = True
except ImportError:
    HAS_PYODBC = False

try:
    from pypdf import PdfReader
    HAS_PDF = True
except ImportError:
    try:
        from PyPDF2 import PdfReader
        HAS_PDF = True
    except ImportError:
        HAS_PDF = False

try:
    import docx as python_docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


TELEGRAM_URL = "https://t.me/Clubapp8"
TELEGRAM_HANDLE = "@Clubapp8"

APP_NAME = "Batch CSV Converter"
APP_VERSION = "2.1.0"
GITHUB_URL = "https://github.com/Mrmmd2004/BatchCSVConverter"

EXCEL_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
LEGACY_EXCEL_EXTS = {".xls"}
CSV_EXTS = {".csv"}
TEXT_EXTS = {".txt", ".log"}
DB_EXTS = {".db", ".sqlite", ".sqlite3", ".db3"}
MDB_EXTS = {".mdb", ".accdb"}
MDF_EXTS = {".mdf"}
SQL_EXTS = {".sql"}
TORRENT_EXTS = {".torrent"}
PDF_EXTS = {".pdf"}
WORD_EXTS = {".docx"}
LEGACY_WORD_EXTS = {".doc"}
SQLITE_MAGIC = b"SQLite format 3\x00"

SAMPLE_BYTES = 65536  # sample size used for encoding/delimiter detection


# ---------------------------------------------------------------------------
# General helpers
# ---------------------------------------------------------------------------
def format_size(num_bytes: float) -> str:
    """Convert a byte count to a readable string with a suitable unit (B/KB/MB/GB/TB)."""
    units = ["B", "KB", "MB", "GB", "TB"]
    size = float(num_bytes)
    for unit in units:
        if size < 1024.0 or unit == units[-1]:
            if unit == "B":
                return f"{int(size)} {unit}"
            return f"{size:.2f} {unit}"
        size /= 1024.0
    return f"{size:.2f} TB"


def is_sqlite_file(path: str) -> bool:
    """Detect a SQLite database from its magic header, regardless of file extension."""
    try:
        with open(path, "rb") as f:
            header = f.read(16)
        return header == SQLITE_MAGIC
    except OSError:
        return False


def looks_binary(chunk: bytes) -> bool:
    """Simple binary-file check based on null bytes or a high share of control characters."""
    if not chunk:
        return False
    if b"\x00" in chunk:
        return True
    text_chars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)) - {0x7f})
    nontext = chunk.translate(None, bytes(text_chars))
    return len(nontext) / max(1, len(chunk)) > 0.30


def detect_encoding(path: str) -> str:
    """Detect a file's encoding. Falls back to a list of common encodings if chardet is missing."""
    try:
        with open(path, "rb") as f:
            raw = f.read(SAMPLE_BYTES)
    except OSError:
        return "utf-8"

    if HAS_CHARDET and raw:
        try:
            result = chardet.detect(raw)
            enc = result.get("encoding")
            if enc:
                return enc
        except Exception:
            pass

    for enc in ("utf-8-sig", "utf-8", "cp1256", "windows-1256", "latin-1"):
        try:
            raw.decode(enc)
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return "latin-1"  # almost always decodes without errors


def detect_dialect(sample_text: str):
    """Detect the text delimiter with csv.Sniffer, with a manual fallback."""
    try:
        dialect = csv.Sniffer().sniff(sample_text, delimiters=",;\t| ")
        return dialect.delimiter
    except csv.Error:
        pass

    # manual fallback: count occurrences of each common delimiter in the first few lines
    lines = [ln for ln in sample_text.splitlines() if ln.strip()][:20]
    if not lines:
        return None

    candidates = ["\t", ",", ";", "|"]
    best_delim, best_score = None, 0
    for delim in candidates:
        counts = [ln.count(delim) for ln in lines]
        if not counts or counts[0] == 0:
            continue
        # a consistent delimiter count across lines = a good delimiter
        consistent = sum(1 for c in counts if c == counts[0])
        score = consistent
        if score > best_score:
            best_score, best_delim = score, delim

    if best_delim:
        return best_delim

    # treat runs of several spaces as a delimiter (column-aligned log formats)
    if any(re.search(r"\s{2,}", ln) for ln in lines):
        return "MULTI_SPACE"

    return None  # no delimiter found -> one column per line


# ---------------------------------------------------------------------------
# Conversion functions per file type
# ---------------------------------------------------------------------------
def convert_excel(path: str, out_path: str):
    ext = os.path.splitext(path)[1].lower()

    if ext in LEGACY_EXCEL_EXTS:
        if not HAS_XLRD:
            raise RuntimeError(
                "The .xls format is legacy and the xlrd library is not installed. "
                "Please run 'pip install xlrd' or convert the file to xlsx."
            )
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            for row_idx in range(sheet.nrows):
                row = sheet.row_values(row_idx)
                writer.writerow(row)
        return

    if not HAS_OPENPYXL:
        raise RuntimeError("The openpyxl library is not installed. Please run 'pip install openpyxl'.")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if v is None else v for v in row])
    wb.close()


def convert_delimited_text(path: str, out_path: str):
    """For CSV / TXT / LOG and any other text file, with automatic delimiter detection."""
    with open(path, "rb") as f:
        raw_head = f.read(SAMPLE_BYTES)

    if looks_binary(raw_head):
        raise RuntimeError("Binary file detected; it cannot be converted to CSV.")

    encoding = detect_encoding(path)

    try:
        with open(path, "r", encoding=encoding, errors="replace", newline="") as f:
            text = f.read()
    except (UnicodeDecodeError, LookupError):
        with open(path, "r", encoding="latin-1", errors="replace", newline="") as f:
            text = f.read()

    sample = text[:SAMPLE_BYTES]
    delimiter = detect_dialect(sample)

    with open(out_path, "w", newline="", encoding="utf-8-sig") as fout:
        writer = csv.writer(fout, quoting=csv.QUOTE_MINIMAL)

        if delimiter is None:
            # no delimiter found: each line = one column
            for line in text.splitlines():
                writer.writerow([line])
        elif delimiter == "MULTI_SPACE":
            for line in text.splitlines():
                if line.strip() == "":
                    writer.writerow([])
                    continue
                fields = re.split(r"\s{2,}", line.strip())
                writer.writerow(fields)
        else:
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            for row in reader:
                writer.writerow(row)


def convert_sqlite(path: str, out_path: str):
    """
    Extract every table of a SQLite database (.db/.sqlite/...) into a single CSV file.
    Each table is marked with a separator line '# Table: <name>', followed by the
    header row and the data rows. A blank line separates tables.
    """
    conn = sqlite3.connect(path)
    try:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT name FROM sqlite_master "
            "WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
        )
        tables = [row[0] for row in cursor.fetchall()]

        if not tables:
            raise RuntimeError("No tables were found in this SQLite database.")

        with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
            for i, table in enumerate(tables):
                if i > 0:
                    writer.writerow([])
                writer.writerow([f"# Table: {table}"])

                table_cursor = conn.execute(f'SELECT * FROM "{table}"')
                col_names = [d[0] for d in table_cursor.description]
                writer.writerow(col_names)
                for row in table_cursor:
                    writer.writerow(row)
    finally:
        conn.close()


def convert_mdb(path: str, out_path: str):
    """Convert a Microsoft Access (.mdb/.accdb) database using pyodbc."""
    if not HAS_PYODBC:
        raise RuntimeError(
            "Converting Access files (.mdb/.accdb) requires the pyodbc library. "
            "Run 'pip install pyodbc' and make sure the 'Microsoft Access Database Engine' "
            "is installed on Windows."
        )
    conn_str = (
        r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ=" + path + ";"
    )
    conn = pyodbc.connect(conn_str, autocommit=True)
    try:
        cursor = conn.cursor()
        tables = [row.table_name for row in cursor.tables(tableType="TABLE")
                  if not row.table_name.startswith("MSys")]

        if not tables:
            raise RuntimeError("No tables were found in this Access database.")

        with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
            for i, table in enumerate(tables):
                if i > 0:
                    writer.writerow([])
                writer.writerow([f"# Table: {table}"])
                tcur = conn.cursor()
                tcur.execute(f"SELECT * FROM [{table}]")
                col_names = [col[0] for col in tcur.description]
                writer.writerow(col_names)
                for row in tcur.fetchall():
                    writer.writerow(list(row))
    finally:
        conn.close()


def convert_mdf(path: str, out_path: str):
    """A .mdf file (SQL Server primary data file) requires a SQL Server engine."""
    raise RuntimeError(
        ".mdf files (SQL Server Data File) cannot be read directly without SQL Server installed. "
        "Please attach the database in SQL Server Management Studio first, "
        "then export it to CSV with the bcp command or the Export Wizard."
    )


# ---------------------------------------------------------------------------
# Fast streaming parser for SQL dumps
# ---------------------------------------------------------------------------
# Instead of walking character by character in a Python loop (very slow for
# multi-gigabyte files), a single regex is used to tokenize the input.
# The regex engine runs in C, so the number of Python-level steps equals the
# number of tokens, not characters, which is several times faster. In addition,
# the whole file is never read into memory at once: it is read in chunks and
# each SQL statement is processed and written to disk as soon as it is complete
# (memory use is bounded by the largest single statement, not the whole file).
_SQL_TOKEN_RE = re.compile(
    r"'(?:[^'\\]|\\.|'')*(?:'|\Z)"      # single-quoted string (or the rest of the buffer if the closing quote has not arrived yet)
    r"|\"(?:[^\"\\]|\\.)*(?:\"|\Z)"     # double-quoted string (same as above)
    r"|[();,]"                          # structural characters
    r"|[^'\"();,]+",                    # any other chunk
    re.DOTALL,
)

_SQL_INSERT_RE = re.compile(
    r"INSERT\s+INTO\s+[`\"\[]?([A-Za-z0-9_]+)[`\"\]]?\s*(\([^)]*\))?\s*VALUES\s*(.*)",
    re.IGNORECASE | re.DOTALL,
)

_SQL_CHUNK_SIZE = 4 * 1024 * 1024  # 4MB


def _iter_sql_statements(f, chunk_size: int = _SQL_CHUNK_SIZE):
    """
    Stream the file and yield each complete SQL statement (up to ';' at depth zero,
    taking strings and nested parentheses into account) without loading the whole file into memory.

    Important implementation note: we never modify the buf string while a finditer is
    live on it (that used to corrupt positions); instead we advance with
    regex.match in a manual loop and trim buf only once, at the start of each new
    chunk read.
    """
    buf = ""
    pos = 0          # current scan position in buf
    stmt_start = 0   # start of the SQL statement being processed, in buf coordinates
    depth = 0        # open-parenthesis depth, preserved across chunks

    while True:
        chunk = f.read(chunk_size)
        eof = not chunk
        if chunk:
            buf += chunk

        while True:
            m = _SQL_TOKEN_RE.match(buf, pos)
            if m is None:
                break  # reached the end of the current data; wait for the next chunk or EOF
            if not eof and m.end() == len(buf):
                # this token may be incomplete (e.g. a string that continues into the next
                # chunk); wait until more data arrives.
                break
            tok = m.group()
            pos = m.end()
            if tok == "(":
                depth += 1
            elif tok == ")":
                depth -= 1
            elif tok == ";" and depth == 0:
                yield buf[stmt_start:pos]
                stmt_start = pos

        if eof:
            if buf[stmt_start:].strip():
                yield buf[stmt_start:]
            return

        # keep memory bounded: discard the fully processed prefix (up to the start of
        # the current statement); this is done only here (between two reads),
        # never in the middle of a finditer scan.
        if stmt_start > 0:
            buf = buf[stmt_start:]
            pos -= stmt_start
            stmt_start = 0


def _split_sql_values(values_str: str):
    """Parse the VALUES part (one or more tuples) of an INSERT statement into a list of rows."""
    rows = []
    fields = None
    depth = 0
    buf_parts = []

    for m in _SQL_TOKEN_RE.finditer(values_str):
        tok = m.group()
        if tok == "(":
            depth += 1
            if depth == 1:
                fields = []
                buf_parts = []
                continue
        elif tok == ")":
            depth -= 1
            if depth == 0 and fields is not None:
                fields.append("".join(buf_parts).strip())
                rows.append(_clean_sql_fields(fields))
                fields = None
                continue
        elif tok == "," and depth == 1:
            fields.append("".join(buf_parts).strip())
            buf_parts = []
            continue

        if fields is not None:
            buf_parts.append(tok)

    return rows


def _clean_sql_fields(fields):
    cleaned = []
    for f in fields:
        f = f.strip()
        if len(f) >= 2 and f[0] == f[-1] and f[0] in ("'", '"'):
            inner = f[1:-1].replace("\\'", "'").replace('\\"', '"').replace("''", "'")
            cleaned.append(inner)
        elif f.upper() == "NULL":
            cleaned.append("")
        else:
            cleaned.append(f)
    return cleaned


def convert_sql_dump(path: str, out_path: str, progress_cb=None):
    """
    Parse an SQL dump file (like mysqldump/pg_dump output) and extract INSERT statements
    to CSV; each table is separated by a '# Table: name' section. If no INSERT is found
    (e.g. schema only), the file is converted as plain text.

    This version is fully streaming: the input file is never loaded into memory at once
    and each table's rows are written to disk immediately (into a temporary file for that
    table). If progress_cb is given, it is called periodically with the progress
    percentage (0..100) so the UI can show status for large files.
    """
    encoding = detect_encoding(path)
    total_size = max(1, os.path.getsize(path))

    tmp_dir = out_path + ".sqlparts"
    os.makedirs(tmp_dir, exist_ok=True)
    table_writers = {}   # table_name -> (file_handle, csv.writer, wrote_header)
    table_order = []
    any_insert_found = False
    fallback_lines_written = False

    def get_writer(table, columns):
        if table not in table_writers:
            fh = open(os.path.join(tmp_dir, f"{len(table_order)}.part"), "w",
                      newline="", encoding="utf-8-sig")
            w = csv.writer(fh, quoting=csv.QUOTE_MINIMAL)
            table_writers[table] = [fh, w, False]
            table_order.append(table)
        entry = table_writers[table]
        if columns and not entry[2]:
            entry[1].writerow(columns)
            entry[2] = True
        return entry[1]

    try:
        with open(path, "r", encoding=encoding, errors="replace") as f:
            last_reported = 0
            stmt_count = 0
            for stmt in _iter_sql_statements(f):
                stmt_count += 1
                m = _SQL_INSERT_RE.search(stmt)
                if m:
                    any_insert_found = True
                    table = m.group(1)
                    col_group = m.group(2)
                    columns = None
                    if col_group:
                        columns = [c.strip(" `\"[]") for c in col_group.strip("()").split(",")]
                    values_part = m.group(3)
                    writer = get_writer(table, columns)
                    for row in _split_sql_values(values_part):
                        writer.writerow(row)

                # report progress every ~256 statements or when the percentage changes noticeably
                if progress_cb and stmt_count % 256 == 0:
                    pos = f.tell()
                    pct = min(99.0, (pos / total_size) * 100.0)
                    if pct - last_reported >= 0.5:
                        last_reported = pct
                        try:
                            progress_cb(pct)
                        except Exception:
                            pass

        if not any_insert_found:
            # no INSERT found; re-read the file line by line and write
            # straight to CSV (without loading the whole file into memory)
            with open(path, "r", encoding=encoding, errors="replace") as f, \
                 open(out_path, "w", newline="", encoding="utf-8-sig") as out_f:
                writer = csv.writer(out_f, quoting=csv.QUOTE_MINIMAL)
                for line in f:
                    writer.writerow([line.rstrip("\n").rstrip("\r")])
            fallback_lines_written = True
        else:
            for fh, _w, _hdr in table_writers.values():
                fh.close()

            with open(out_path, "w", newline="", encoding="utf-8-sig") as out_f:
                writer = csv.writer(out_f, quoting=csv.QUOTE_MINIMAL)
                for i, table in enumerate(table_order):
                    if i > 0:
                        writer.writerow([])
                    writer.writerow([f"# Table: {table}"])
                    part_path = os.path.join(tmp_dir, f"{i}.part")
                    with open(part_path, "r", encoding="utf-8-sig", newline="") as part_f:
                        for row in csv.reader(part_f):
                            writer.writerow(row)
    finally:
        if not fallback_lines_written:
            for fh, _w, _hdr in table_writers.values():
                try:
                    fh.close()
                except Exception:
                    pass
        for i in range(len(table_order)):
            try:
                os.remove(os.path.join(tmp_dir, f"{i}.part"))
            except OSError:
                pass
        try:
            os.rmdir(tmp_dir)
        except OSError:
            pass

    if progress_cb:
        try:
            progress_cb(100.0)
        except Exception:
            pass


def _bdecode(data: bytes):
    """Simple bencode decoder for .torrent files (no external dependencies)."""
    pos = 0

    def decode():
        nonlocal pos
        c = data[pos:pos + 1]
        if c == b"i":
            pos += 1
            end = data.index(b"e", pos)
            val = int(data[pos:end])
            pos = end + 1
            return val
        if c == b"l":
            pos += 1
            lst = []
            while data[pos:pos + 1] != b"e":
                lst.append(decode())
            pos += 1
            return lst
        if c == b"d":
            pos += 1
            d = {}
            while data[pos:pos + 1] != b"e":
                key = decode()
                d[key] = decode()
            pos += 1
            return d
        if c.isdigit():
            colon = data.index(b":", pos)
            length = int(data[pos:colon])
            start = colon + 1
            val = data[start:start + length]
            pos = start + length
            return val
        raise ValueError(f"Invalid bencode at position {pos}")

    return decode()


def convert_torrent(path: str, out_path: str):
    """Extract .torrent metadata (name, trackers, file list) into CSV."""
    with open(path, "rb") as f:
        data = f.read()

    try:
        torrent = _bdecode(data)
    except Exception as exc:
        raise RuntimeError(f"Invalid torrent file or it could not be decoded: {exc}")

    def s(b):
        if isinstance(b, bytes):
            try:
                return b.decode("utf-8")
            except UnicodeDecodeError:
                return b.decode("latin-1", errors="replace")
        return b

    info = torrent.get(b"info", {}) if isinstance(torrent, dict) else {}
    trackers = []
    if b"announce-list" in torrent:
        for group in torrent[b"announce-list"]:
            for t in group:
                trackers.append(s(t))
    elif b"announce" in torrent:
        trackers.append(s(torrent[b"announce"]))

    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(["# Torrent Info"])
        writer.writerow(["Key", "Value"])
        writer.writerow(["Name", s(info.get(b"name", b""))])
        writer.writerow(["Comment", s(torrent.get(b"comment", b""))])
        writer.writerow(["Created By", s(torrent.get(b"created by", b""))])
        writer.writerow(["Creation Date", torrent.get(b"creation date", "")])
        writer.writerow(["Piece Length", info.get(b"piece length", "")])

        if trackers:
            writer.writerow([])
            writer.writerow(["# Trackers"])
            for t in trackers:
                writer.writerow([t])

        writer.writerow([])
        writer.writerow(["# Files"])
        writer.writerow(["Path", "Size (bytes)"])
        if b"files" in info:
            for entry in info[b"files"]:
                parts = [s(p) for p in entry.get(b"path", [])]
                writer.writerow(["/".join(parts), entry.get(b"length", "")])
        else:
            writer.writerow([s(info.get(b"name", b"")), info.get(b"length", "")])


def convert_pdf(path: str, out_path: str):
    """Extract the text of each PDF page and convert it line by line to CSV (with delimiter detection)."""
    if not HAS_PDF:
        raise RuntimeError("The pypdf library is not installed, so PDF conversion is unavailable. Please run 'pip install pypdf'.")

    reader = PdfReader(path)
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            writer.writerow([f"# Page {i + 1}"])
            for line in text.splitlines():
                if line.strip() == "":
                    continue
                delim = detect_dialect(line)
                if delim == "MULTI_SPACE":
                    row = re.split(r"\s{2,}", line.strip())
                elif delim:
                    row = next(csv.reader([line], delimiter=delim))
                else:
                    row = [line]
                writer.writerow(row)
            writer.writerow([])


def convert_word(path: str, out_path: str):
    """Extract paragraphs and tables of a Word (.docx) file into CSV."""
    ext = os.path.splitext(path)[1].lower()
    if ext in LEGACY_WORD_EXTS:
        raise RuntimeError(
            "Legacy .doc files (Word 97-2003) are not supported directly. "
            "Please open the file in Word and use Save As to save it as .docx."
        )
    if not HAS_DOCX:
        raise RuntimeError("Converting Word (.docx) files requires the python-docx library. Run 'pip install python-docx'.")

    document = python_docx.Document(path)
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        wrote = False

        paragraphs = [p.text for p in document.paragraphs if p.text.strip() != ""]
        if paragraphs:
            writer.writerow(["# Paragraphs"])
            for p in paragraphs:
                writer.writerow([p])
            wrote = True

        for ti, table in enumerate(document.tables):
            if wrote:
                writer.writerow([])
            writer.writerow([f"# Table {ti + 1}"])
            for row in table.rows:
                writer.writerow([cell.text for cell in row.cells])
            wrote = True

        if not wrote:
            writer.writerow(["(The document is empty)"])


def plan_output_paths(files, out_dir=None):
    """Decide the output .csv path of every file before the run starts.

    out_dir=None -> next to each source file; otherwise everything goes into out_dir.
    If two sources would produce the same name inside one run (a.xlsx and a.txt, or the same
    name from two different folders), the later ones become "a (2).csv", "a (3).csv", ...
    so that no output overwrites another.
    """
    plan, used = {}, set()
    for p in files:
        folder = out_dir or os.path.dirname(p)
        base = os.path.splitext(os.path.basename(p))[0]
        n = 1
        while True:
            cand = os.path.join(folder, base + (".csv" if n == 1 else f" ({n}).csv"))
            key = os.path.normcase(os.path.abspath(cand))
            if key not in used:
                break
            n += 1
        used.add(key)
        plan[p] = cand
    return plan


def convert_file(path: str, progress_queue=None, out_path=None):
    """
    Convert one file to CSV based on its extension.
    By default the output is saved in the same folder as the source file, with the same name and
    a .csv extension. If out_path is given, the CSV is written there (folders are created).
    Returns: (success: bool, out_path_or_None, error_message_or_None)

    progress_queue: if given (a multiprocessing.Manager().Queue()), ("sql_progress", path, percent)
    messages are put on it for large SQL files so the UI can show progress within that file.
    Each file runs in a separate process, so the queue must be a multiprocessing queue,
    not a regular queue.Queue.
    """
    try:
        if out_path is None:
            folder = os.path.dirname(path)
            base = os.path.splitext(os.path.basename(path))[0]
            out_path = os.path.join(folder, base + ".csv")
        else:
            os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)

        # if the output has the same name as the input (a .csv file that is its own input),
        # avoid overwriting it in place.
        same_file = os.path.normcase(os.path.abspath(out_path)) == os.path.normcase(os.path.abspath(path))
        real_out_path = out_path + ".tmp" if same_file else out_path

        ext = os.path.splitext(path)[1].lower()

        if ext in EXCEL_EXTS or ext in LEGACY_EXCEL_EXTS:
            convert_excel(path, real_out_path)
        elif ext in DB_EXTS or is_sqlite_file(path):
            convert_sqlite(path, real_out_path)
        elif ext in MDB_EXTS:
            convert_mdb(path, real_out_path)
        elif ext in MDF_EXTS:
            convert_mdf(path, real_out_path)
        elif ext in SQL_EXTS:
            cb = None
            if progress_queue is not None:
                cb = lambda pct: progress_queue.put(("sql_progress", path, pct))
            convert_sql_dump(path, real_out_path, progress_cb=cb)
        elif ext in TORRENT_EXTS:
            convert_torrent(path, real_out_path)
        elif ext in PDF_EXTS:
            convert_pdf(path, real_out_path)
        elif ext in WORD_EXTS or ext in LEGACY_WORD_EXTS:
            convert_word(path, real_out_path)
        else:
            # CSV, TXT, LOG and any other unknown extension (or none) go through this path
            convert_delimited_text(path, real_out_path)

        if same_file:
            os.replace(real_out_path, out_path)

        return True, out_path, None

    except Exception as exc:  # any error on one file must not block the other files
        return False, None, f"{type(exc).__name__}: {exc}"


# ---------------------------------------------------------------------------
# Shared batch conversion engine (no dependency on the GUI)
# ---------------------------------------------------------------------------
def run_batch_conversion(files, on_sql_progress=None, on_file_done=None, out_dir=None):
    """
    Run batch conversion of files with a ProcessPoolExecutor (true parallelism across CPU
    cores). This function has no dependency on Tkinter and can be called both from the GUI
    (in a background thread) and from the headless command-line mode.

    on_sql_progress(path, percent): called during processing of large SQL files.
    on_file_done(path, success, out_path, error, processed_count, total): called after each file.
    out_dir: folder for all output CSVs; None = next to each source file.
    """
    total = len(files)
    plan = plan_output_paths(files, out_dir)
    max_workers = min(max(1, os.cpu_count() or 4), total) if total else 1

    manager = multiprocessing.Manager()
    progress_queue = manager.Queue()
    processed = 0

    try:
        with ProcessPoolExecutor(max_workers=max_workers) as executor:
            future_to_path = {
                executor.submit(convert_file, p, progress_queue, plan[p]): p for p in files
            }
            pending = set(future_to_path)

            while pending:
                try:
                    while True:
                        _kind, path, pct = progress_queue.get_nowait()
                        if on_sql_progress:
                            on_sql_progress(path, pct)
                except queue.Empty:
                    pass

                done_now, pending = cf_wait(pending, timeout=0.1, return_when=cf_FIRST_COMPLETED)
                for future in done_now:
                    path = future_to_path[future]
                    try:
                        success, out_path, error = future.result()
                    except Exception as exc:  # extra safety against unexpected errors
                        success, out_path, error = False, None, f"Unexpected: {exc}"
                    processed += 1
                    if on_file_done:
                        on_file_done(path, success, out_path, error, processed, total)
    finally:
        manager.shutdown()


def _run_cli_mode(paths, out_dir=None):
    """
    Lightweight headless mode: no Tkinter, progress is just printed to the console.
    Run: python batch_csv_converter.py file1.sql file2.xlsx ...
    No graphical overhead; suited for scripting and maximum speed.
    """
    files = [p for p in paths if os.path.isfile(p)]
    missing = [p for p in paths if not os.path.isfile(p)]
    for p in missing:
        print(f"[!] Not found, skipped: {p}")

    if not files:
        print("No valid files were given for conversion.")
        return

    total = len(files)
    print(f"Starting conversion of {total} file(s) with {min(max(1, os.cpu_count() or 4), total)} parallel process(es)...")

    failed = []

    def on_progress(path, pct):
        print(f"\r  {os.path.basename(path)}: {pct:5.1f}%", end="", flush=True)

    def on_done(path, success, out_path, error, done_count, total_count):
        print(f"\r[{done_count}/{total_count}] {'OK ' if success else 'FAIL'}  {os.path.basename(path)}"
              + (f"  ->  {os.path.basename(out_path)}" if success else f"  Error: {error}"))
        if not success:
            failed.append((path, error))

    run_batch_conversion(files, on_sql_progress=on_progress, on_file_done=on_done, out_dir=out_dir)

    print(f"\nDone: {total - len(failed)} succeeded / {total} total.")
    if failed:
        print("Failed files:")
        for p, err in failed:
            print(f"  - {os.path.basename(p)}: {err}")


# ---------------------------------------------------------------------------
# Modern GUI (Tkinter + ttk, no extra dependencies)
# ---------------------------------------------------------------------------
import json
import subprocess
import ctypes

HAS_DND = os.name == "nt"  # native drag & drop (WM_DROPFILES), no extra packages needed


class _DropTarget:
    """Native Windows drag & drop of files using only ctypes (no third-party packages).

    Uses the official comctl32 window-subclassing API (SetWindowSubclass / DefSubclassProc),
    which chains to the original window procedure safely. The native callback only copies
    the dropped paths and hands them over; all real work happens later on the Tk event loop.
    """
    WM_DROPFILES = 0x0233
    WM_COPYGLOBALDATA = 0x0049
    WM_NCDESTROY = 0x0082

    def __init__(self, tk_root, on_files):
        from ctypes import wintypes
        self._on_files = on_files
        self._keep = []  # keep callbacks alive (otherwise they get garbage collected)
        self._user32 = user32 = ctypes.WinDLL("user32")
        self._shell32 = shell32 = ctypes.WinDLL("shell32")
        self._comctl32 = comctl32 = ctypes.WinDLL("comctl32")

        lresult = ctypes.c_ssize_t
        self._proc_type = ctypes.WINFUNCTYPE(
            lresult, wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM,
            ctypes.c_size_t, ctypes.c_size_t)
        comctl32.SetWindowSubclass.argtypes = [
            wintypes.HWND, self._proc_type, ctypes.c_size_t, ctypes.c_size_t]
        comctl32.SetWindowSubclass.restype = wintypes.BOOL
        comctl32.RemoveWindowSubclass.argtypes = [
            wintypes.HWND, self._proc_type, ctypes.c_size_t]
        comctl32.RemoveWindowSubclass.restype = wintypes.BOOL
        comctl32.DefSubclassProc.argtypes = [
            wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
        comctl32.DefSubclassProc.restype = lresult

        user32.GetParent.restype = wintypes.HWND
        user32.GetParent.argtypes = [wintypes.HWND]
        shell32.DragAcceptFiles.argtypes = [wintypes.HWND, wintypes.BOOL]
        shell32.DragQueryFileW.restype = wintypes.UINT
        shell32.DragQueryFileW.argtypes = [ctypes.c_void_p, wintypes.UINT,
                                           wintypes.LPWSTR, wintypes.UINT]
        shell32.DragFinish.argtypes = [ctypes.c_void_p]
        try:
            user32.ChangeWindowMessageFilterEx.argtypes = [
                wintypes.HWND, wintypes.UINT, wintypes.DWORD, ctypes.c_void_p]
        except Exception:
            pass

        tk_root.update_idletasks()
        inner = tk_root.winfo_id()
        outer = user32.GetParent(inner)
        hooked = 0
        for hwnd in {h for h in (inner, outer) if h}:
            if self._hook(hwnd):
                hooked += 1
        if not hooked:
            raise OSError("could not hook the window for drag & drop")

    def _hook(self, hwnd):
        comctl32 = self._comctl32

        def proc(h, msg, wparam, lparam, uid, ref):
            try:
                if msg == self.WM_DROPFILES:
                    paths = self._read_drop(wparam)
                    if paths:
                        self._on_files(paths)
                    return 0
                if msg == self.WM_NCDESTROY:
                    comctl32.RemoveWindowSubclass(h, callback, uid)
            except Exception:
                if msg == self.WM_DROPFILES:
                    return 0
            return comctl32.DefSubclassProc(h, msg, wparam, lparam)

        callback = self._proc_type(proc)
        if not comctl32.SetWindowSubclass(hwnd, callback, 1, 0):
            return False
        self._keep.append(callback)
        self._shell32.DragAcceptFiles(hwnd, True)
        try:  # allow drops even if the app runs elevated (UAC message filter)
            for msg in (self.WM_DROPFILES, self.WM_COPYGLOBALDATA):
                self._user32.ChangeWindowMessageFilterEx(hwnd, msg, 1, None)
        except Exception:
            pass
        return True

    def _read_drop(self, hdrop):
        shell32 = self._shell32
        paths = []
        try:
            count = shell32.DragQueryFileW(hdrop, 0xFFFFFFFF, None, 0)
            for i in range(count):
                length = shell32.DragQueryFileW(hdrop, i, None, 0)
                buf = ctypes.create_unicode_buffer(length + 1)
                shell32.DragQueryFileW(hdrop, i, buf, length + 1)
                paths.append(buf.value)
        finally:
            shell32.DragFinish(hdrop)
        return paths


THEMES = {
    "dark": {
        "bg": "#0f1115", "surface": "#171a21", "surface2": "#1f232d", "border": "#2b303c",
        "fg": "#e7e9ee", "dim": "#8b93a3", "accent": "#6c8cff", "accent_hover": "#5a7af0",
        "accent_fg": "#ffffff", "ok": "#3ecf8e", "err": "#ff6b6b", "warn": "#f5b94a",
    },
    "light": {
        "bg": "#f3f5fa", "surface": "#ffffff", "surface2": "#edf0f7", "border": "#dce1ec",
        "fg": "#1b1f2a", "dim": "#6a7385", "accent": "#4f6df5", "accent_hover": "#3d59dc",
        "accent_fg": "#ffffff", "ok": "#15996a", "err": "#d93f49", "warn": "#b9790f",
    },
}

FONT = "Segoe UI"


def make_circle_png_b64(hex_color: str, size: int) -> str:
    """Anti-aliased filled circle as a base64 PNG (transparent background, no extra packages)."""
    import zlib, struct, base64
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    ss = 4                                   # supersampling for smooth edges
    c = size / 2.0
    rad = size / 2.0 - 0.5
    raw = bytearray()
    for y in range(size):
        raw.append(0)                        # PNG filter type 0
        for x in range(size):
            hits = 0
            for sy in range(ss):
                for sx in range(ss):
                    px = x + (sx + 0.5) / ss - c
                    py = y + (sy + 0.5) / ss - c
                    if px * px + py * py <= rad * rad:
                        hits += 1
            raw += bytes((r, g, b, int(255 * hits / (ss * ss))))

    def chunk(tag, data):
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    png = (b"\x89PNG\r\n\x1a\n"
           + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 6, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(bytes(raw), 9))
           + chunk(b"IEND", b""))
    return base64.b64encode(png).decode("ascii")


ON_COLOR = "#22c55e"    # green circle = file is ON (will be converted)
OFF_COLOR = "#ef4444"   # red circle   = file is OFF (skipped)

_LOG_FILE = None


def _log_path() -> str:
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    return os.path.join(base, "BatchCsvConverter", "crash.log")


def _log_exception(where: str = "") -> None:
    """Write the current exception to crash.log (the GUI has no console in the exe)."""
    try:
        path = _log_path()
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "a", encoding="utf-8") as f:
            f.write(f"--- {where} ---\n{traceback.format_exc()}\n")
    except Exception:
        pass


def _enable_crash_log() -> None:
    """Log Python errors and hard crashes (faulthandler) to %APPDATA%\\BatchCsvConverter\\crash.log."""
    global _LOG_FILE
    try:
        import faulthandler
        path = _log_path()
        os.makedirs(os.path.dirname(path), exist_ok=True)
        _LOG_FILE = open(path, "a", buffering=1)
        faulthandler.enable(_LOG_FILE)
    except Exception:
        pass


def resource_path(name: str) -> str:
    """Path of bundled files; works both as a script and inside the exe (PyInstaller)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


def _settings_path() -> str:
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    return os.path.join(base, "BatchCsvConverter", "settings.json")


def load_settings() -> dict:
    try:
        with open(_settings_path(), "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_settings(data: dict) -> None:
    try:
        path = _settings_path()
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f)
    except Exception:
        pass


def file_kind(path: str) -> str:
    """Readable file type name (display only, used in the table)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in EXCEL_EXTS or ext in LEGACY_EXCEL_EXTS:
        return "Excel"
    if ext in CSV_EXTS:
        return "CSV"
    if ext in TEXT_EXTS:
        return "Text"
    if ext in MDB_EXTS:
        return "Access"
    if ext in MDF_EXTS:
        return "SQL Server"
    if ext in SQL_EXTS:
        return "SQL dump"
    if ext in TORRENT_EXTS:
        return "Torrent"
    if ext in PDF_EXTS:
        return "PDF"
    if ext in WORD_EXTS or ext in LEGACY_WORD_EXTS:
        return "Word"
    if ext in DB_EXTS or is_sqlite_file(path):
        return "SQLite"
    return "Other"


def open_in_explorer(path: str) -> None:
    """Open a folder (and select the file when possible)."""
    try:
        if os.name == "nt":
            if os.path.isfile(path):
                subprocess.Popen(["explorer", f"/select,{os.path.normpath(path)}"])
            else:
                os.startfile(path)  # type: ignore[attr-defined]
        else:
            webbrowser.open(os.path.dirname(path) if os.path.isfile(path) else path)
    except Exception:
        pass


class FlatButton(tk.Label):
    """Flat, modern button with hover effect (replacement for tk.Button)."""

    def __init__(self, parent, text, command, kind="ghost", font=None, padx=16, pady=8):
        super().__init__(parent, text=text, cursor="hand2", padx=padx, pady=pady,
                         font=font or (FONT, 10))
        self.command = command
        self.kind = kind
        self._enabled = True
        self._hover = False
        self._pal = None
        self.bind("<Enter>", self._on_enter)
        self.bind("<Leave>", self._on_leave)
        self.bind("<Button-1>", self._on_click)

    def set_palette(self, pal):
        self._pal = pal
        self._paint()

    def set_enabled(self, enabled: bool):
        self._enabled = enabled
        self.config(cursor="hand2" if enabled else "arrow")
        self._paint()

    def _paint(self):
        p = self._pal
        if not p:
            return
        hot = self._hover and self._enabled
        if self.kind == "primary":
            bg = p["accent_hover"] if hot else p["accent"]
            fg = p["accent_fg"]
        else:
            bg = p["border"] if hot else p["surface2"]
            fg = p["fg"]
        if not self._enabled:
            bg, fg = p["surface2"], p["dim"]
        self.config(bg=bg, fg=fg)

    def _on_enter(self, _e):
        self._hover = True
        self._paint()

    def _on_leave(self, _e):
        self._hover = False
        self._paint()

    def _on_click(self, _e):
        if self._enabled and self.command:
            self.command()


class BatchCsvConverterApp:
    COLUMNS = (
        ("name", "Name", 200, "w", True),
        ("folder", "Folder", 170, "w", True),
        ("kind", "Type", 80, "center", False),
        ("size", "Size", 80, "e", False),
        ("status", "Status", 120, "w", False),
    )

    def __init__(self, root: tk.Tk, initial_files=None):
        self.root = root
        self.scale = max(1.0, root.winfo_fpixels("1i") / 96.0)

        root.title(f"{APP_NAME}")
        try:
            root.iconbitmap(resource_path("clubapp.ico"))
        except Exception:
            pass

        self.settings = load_settings()
        self.theme_name = self.settings.get("theme", "dark")
        if self.theme_name not in THEMES:
            self.theme_name = "dark"

        self.out_mode = self.settings.get("out_mode", "source")   # "source" or "custom"
        if self.out_mode not in ("source", "custom"):
            self.out_mode = "source"
        self.out_dir = self.settings.get("out_dir", "") or ""
        self._run_out_dir = None

        # state
        self.rows = {}            # iid -> dict(path, size)
        self.path_to_iid = {}     # path -> iid
        self.results = {}         # iid -> (success, out_path, error)
        self.partial = {}         # iid -> progress fraction of SQL files being processed
        self._iid_counter = 0
        self.processed_files = 0
        self.success_count = 0
        self.total_active = 0     # number of files switched ON in the current run
        self.is_processing = False
        self.last_output = None
        self.ui_queue = queue.Queue()

        self._themed = []         # (widget, {option: palette_key})
        self._buttons = []

        self.style = ttk.Style()
        try:
            self.style.theme_use("clam")
        except tk.TclError:
            pass
        self.style.layout("Treeview", [("Treeview.treearea", {"sticky": "nswe"})])
        self.style.layout("Modern.Vertical.TScrollbar", [
            ("Vertical.Scrollbar.trough", {"sticky": "ns", "children": [
                ("Vertical.Scrollbar.thumb", {"expand": "1", "sticky": "nswe"})]})
        ])

        self._build_ui()
        self.apply_theme()
        self._refresh_output_ui()
        self._fit_window()
        self._setup_dnd()

        root.bind("<Control-o>", lambda e: self.on_add_files())
        root.bind("<Delete>", lambda e: self.on_remove_selected())
        root.bind("<Control-a>", lambda e: self.on_select_all())

        if initial_files:
            self.add_paths(initial_files)

        self._poll_queue()

    # ------------------------------------------------------------------
    # UI helpers
    # ------------------------------------------------------------------
    def sc(self, v):
        return int(round(v * self.scale))

    def _frame(self, parent, bg="bg", **kw):
        f = tk.Frame(parent, **kw)
        self._themed.append((f, {"bg": bg}))
        return f

    def _label(self, parent, text, fg="fg", bg="bg", font=None, **kw):
        lb = tk.Label(parent, text=text, font=font or (FONT, 10), **kw)
        self._themed.append((lb, {"fg": fg, "bg": bg}))
        return lb

    def _button(self, parent, text, command, kind="ghost", font=None, padx=16, pady=8):
        b = FlatButton(parent, text, command, kind=kind, font=font,
                       padx=self.sc(padx), pady=self.sc(pady))
        self._buttons.append(b)
        return b

    def _link(self, parent, text, url, bg="bg"):
        lb = tk.Label(parent, text=text, cursor="hand2", font=(FONT, 9, "underline"))
        self._themed.append((lb, {"fg": "accent", "bg": bg}))
        lb.bind("<Button-1>", lambda e: webbrowser.open(url))
        return lb

    # ------------------------------------------------------------------
    # Building the UI
    # ------------------------------------------------------------------
    def _fit_window(self):
        """Size the window to the screen: never bigger than the screen, never smaller
        than what is needed to show every control (so the buttons are always visible)."""
        root = self.root
        root.update_idletasks()
        sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
        req_w, req_h = root.winfo_reqwidth(), root.winfo_reqheight()
        max_w, max_h = int(sw * 0.94), int(sh * 0.90)   # leave room for taskbar / title bar
        min_w = min(max(req_w, self.sc(640)), max_w)
        min_h = min(req_h, max_h)
        w = min(max(self.sc(920), min_w), max_w)
        h = min(max(self.sc(660), min_h), max_h)
        x = max(0, (sw - w) // 2)
        y = max(0, (sh - h) // 3)
        root.minsize(min_w, min_h)
        root.geometry(f"{w}x{h}+{x}+{y}")

    def _build_ui(self):
        pad = self.sc(22)
        root = self.root

        # --- header ---
        header = self._frame(root)
        header.pack(side="top", fill="x", padx=pad, pady=(pad, self.sc(12)))

        left = self._frame(header)
        left.pack(side="left")
        self._label(left, APP_NAME, font=(FONT, 20, "bold")).pack(anchor="w")
        self._label(left, "Convert many files to CSV in bulk.",
                    fg="dim", font=(FONT, 9)).pack(anchor="w", pady=(2, 0))

        right = self._frame(header)
        right.pack(side="right")
        self.theme_btn = self._button(right, "", self.toggle_theme, padx=12, pady=6)
        self.theme_btn.pack(side="right", padx=(8, 0))
        self._button(right, "Telegram", lambda: webbrowser.open(TELEGRAM_URL),
                     padx=12, pady=6).pack(side="right", padx=(8, 0))
        self._button(right, "GitHub", lambda: webbrowser.open(GITHUB_URL),
                     padx=12, pady=6).pack(side="right")

        # --- bottom bars: packed BEFORE the file list (side="bottom") so that, when the
        #     window is small, the file list shrinks instead of the buttons being cut off ---
        footer = self._frame(root)
        footer.pack(side="bottom", fill="x", padx=pad, pady=(0, self.sc(12)))
        self._label(footer, f"v{APP_VERSION}", fg="dim", font=(FONT, 9)).pack(side="left")
        self._link(footer, "GitHub", GITHUB_URL).pack(side="right")
        self._label(footer, "·", fg="dim", font=(FONT, 9)).pack(side="right", padx=6)
        self._link(footer, TELEGRAM_HANDLE, TELEGRAM_URL).pack(side="right")

        actions = self._frame(root)
        actions.pack(side="bottom", fill="x", padx=pad, pady=(self.sc(12), self.sc(8)))
        self.start_btn = self._button(actions, "Start Conversion", self.on_start,
                                      kind="primary", font=(FONT, 11, "bold"), padx=28, pady=10)
        self.start_btn.pack(side="left")
        self.open_btn = self._button(actions, "Open Output Folder", self.on_open_output, pady=10)
        self.open_btn.pack(side="left", padx=(10, 0))
        self.start_btn.set_enabled(False)
        self.open_btn.set_enabled(False)

        prog = self._frame(root)
        prog.pack(side="bottom", fill="x", padx=pad, pady=(self.sc(12), 0))
        top = self._frame(prog)
        top.pack(fill="x")
        self.status_label = self._label(top, "Ready.", fg="dim", font=(FONT, 9), anchor="w")
        self.status_label.pack(side="left")
        self.percent_label = self._label(top, "0%", font=(FONT, 10, "bold"))
        self.percent_label.pack(side="right")
        self.progress = ttk.Progressbar(prog, mode="determinate", maximum=100,
                                        style="Modern.Horizontal.TProgressbar")
        self.progress.pack(fill="x", pady=(self.sc(8), 0))

        # --- where the CSV files are saved ---
        outbar = self._frame(root)
        outbar.pack(side="bottom", fill="x", padx=pad, pady=(self.sc(12), 0))
        self._label(outbar, "Save CSV files to:", fg="dim", font=(FONT, 9)).pack(
            side="left", padx=(0, 10))
        self.out_src_btn = self._button(outbar, "Next to each file", self.on_out_source,
                                        padx=12, pady=6)
        self.out_src_btn.pack(side="left")
        self.out_dir_btn = self._button(outbar, "A folder I choose\u2026", self.on_out_choose,
                                        padx=12, pady=6)
        self.out_dir_btn.pack(side="left", padx=(8, 0))
        self.out_label = self._label(outbar, "", fg="dim", font=(FONT, 9), anchor="w")
        self.out_label.pack(side="left", padx=(12, 0))
        self.out_btns = [self.out_src_btn, self.out_dir_btn]

        # --- main card (takes all the remaining space) ---
        card = tk.Frame(root, highlightthickness=1)
        self._themed.append((card, {"bg": "surface", "highlightbackground": "border",
                                    "highlightcolor": "border"}))
        card.pack(side="top", fill="both", expand=True, padx=pad)

        toolbar = self._frame(card, bg="surface")
        toolbar.pack(side="top", fill="x", padx=self.sc(12), pady=self.sc(10))
        self._button(toolbar, "+  Add Files", self.on_add_files, kind="primary",
                     font=(FONT, 10, "bold")).pack(side="left")
        self._button(toolbar, "Add Folder", self.on_add_folder).pack(side="left", padx=(8, 0))
        self.select_all_btn = self._button(toolbar, "Select All", self.on_select_all)
        self.select_all_btn.pack(side="left", padx=(16, 0))
        self.deselect_all_btn = self._button(toolbar, "Deselect All", self.on_deselect_all)
        self.deselect_all_btn.pack(side="left", padx=(8, 0))
        self.clear_btn = self._button(toolbar, "Remove All", self.on_clear)
        self.clear_btn.pack(side="right")
        self.remove_btn = self._button(toolbar, "Remove Selected", self.on_remove_selected)
        self.remove_btn.pack(side="right", padx=(0, 8))

        self.summary_label = self._label(card, "0 files  ·  0 B", fg="dim", bg="surface",
                                         font=(FONT, 9), anchor="w")
        self.summary_label.pack(side="bottom", fill="x", padx=self.sc(14), pady=self.sc(8))

        tree_wrap = self._frame(card, bg="surface")
        tree_wrap.pack(side="top", fill="both", expand=True, padx=self.sc(1), pady=0)

        self.tree = ttk.Treeview(tree_wrap, columns=[c[0] for c in self.COLUMNS],
                                 show="tree headings", selectmode="extended", height=4)
        circle = self.sc(16)
        self.img_on = tk.PhotoImage(data=make_circle_png_b64(ON_COLOR, circle))
        self.img_off = tk.PhotoImage(data=make_circle_png_b64(OFF_COLOR, circle))
        self.tree.column("#0", width=self.sc(92), minwidth=self.sc(92), stretch=False, anchor="center")
        self.tree.heading("#0", text="Convert", anchor="center")
        for key, title, width, anchor, stretch in self.COLUMNS:
            self.tree.heading(key, text=title, anchor="w" if anchor == "w" else anchor)
            self.tree.column(key, width=self.sc(width), anchor=anchor, stretch=stretch,
                             minwidth=self.sc(60))
        vsb = ttk.Scrollbar(tree_wrap, orient="vertical", command=self.tree.yview,
                            style="Modern.Vertical.TScrollbar")
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")
        self.tree.bind("<Double-1>", self.on_row_double_click)
        self.tree.bind("<Button-1>", self._on_tree_click)
        self.tree.bind("<space>", self._on_tree_space)

        hint = ("Drag & drop files or folders here\nor click \u201cAdd Files\u201d / \u201cAdd Folder\u201d"
                if HAS_DND else "Click \u201cAdd Files\u201d or \u201cAdd Folder\u201d to begin.")
        self.empty_label = tk.Label(tree_wrap, text=hint, font=(FONT, 11), justify="center")
        self._themed.append((self.empty_label, {"fg": "dim", "bg": "surface"}))
        self.empty_label.place(relx=0.5, rely=0.5, anchor="center")

    # ------------------------------------------------------------------
    # Theme
    # ------------------------------------------------------------------
    def apply_theme(self):
        p = THEMES[self.theme_name]
        self.root.configure(bg=p["bg"])
        for widget, spec in self._themed:
            try:
                widget.configure(**{opt: p[key] for opt, key in spec.items()})
            except tk.TclError:
                pass
        for b in self._buttons:
            b.set_palette(p)
        self.theme_btn.config(text="☀  Light" if self.theme_name == "dark" else "☾  Dark")

        st = self.style
        st.configure("Treeview", background=p["surface"], fieldbackground=p["surface"],
                     foreground=p["fg"], borderwidth=0, rowheight=self.sc(34), font=(FONT, 10))
        st.map("Treeview", background=[("selected", p["accent"])],
               foreground=[("selected", p["accent_fg"])])
        st.configure("Treeview.Heading", background=p["surface2"], foreground=p["dim"],
                     relief="flat", borderwidth=0, font=(FONT, 9, "bold"),
                     padding=(self.sc(10), self.sc(7)))
        st.map("Treeview.Heading", background=[("active", p["surface2"])])
        st.configure("Modern.Horizontal.TProgressbar", troughcolor=p["surface2"],
                     background=p["accent"], bordercolor=p["surface2"],
                     lightcolor=p["accent"], darkcolor=p["accent"], thickness=self.sc(8))
        st.configure("Modern.Vertical.TScrollbar", troughcolor=p["surface"],
                     background=p["border"], bordercolor=p["surface"],
                     lightcolor=p["border"], darkcolor=p["border"], relief="flat",
                     gripcount=0, width=self.sc(10))
        st.map("Modern.Vertical.TScrollbar", background=[("active", p["dim"])])

        self.tree.tag_configure("ready", foreground=p["fg"])
        self.tree.tag_configure("queued", foreground=p["dim"])
        self.tree.tag_configure("off", foreground=p["dim"])
        self.tree.tag_configure("working", foreground=p["accent"])
        self.tree.tag_configure("done", foreground=p["ok"])
        self.tree.tag_configure("failed", foreground=p["err"])
        self._set_titlebar_dark(self.theme_name == "dark")

    def toggle_theme(self):
        self.theme_name = "light" if self.theme_name == "dark" else "dark"
        self.settings["theme"] = self.theme_name
        save_settings(self.settings)
        self.apply_theme()

    def _set_titlebar_dark(self, dark: bool):
        """Dark title bar on Windows 10/11 (when supported; otherwise ignored)."""
        if os.name != "nt":
            return
        try:
            self.root.update_idletasks()
            hwnd = ctypes.windll.user32.GetParent(self.root.winfo_id())
            value = ctypes.c_int(1 if dark else 0)
            for attr in (20, 19):  # DWMWA_USE_IMMERSIVE_DARK_MODE (new / old)
                ctypes.windll.dwmapi.DwmSetWindowAttribute(
                    hwnd, attr, ctypes.byref(value), ctypes.sizeof(value))
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Drag & Drop (optional)
    # ------------------------------------------------------------------
    def _setup_dnd(self):
        self._drop = None
        if not HAS_DND:
            return
        try:
            # the native callback only puts the paths in a queue; _poll_queue handles them
            self._drop = _DropTarget(
                self.root, lambda paths: self.ui_queue.put({"type": "drop", "paths": list(paths)}))
        except Exception:
            self._drop = None  # drag & drop is a convenience; the app works without it

    # ------------------------------------------------------------------
    # File list management
    # ------------------------------------------------------------------
    def add_paths(self, paths):
        if self.is_processing:
            return
        expanded = []
        for p in paths:
            p = os.path.abspath(p)
            if os.path.isdir(p):
                try:
                    for name in sorted(os.listdir(p)):
                        fp = os.path.join(p, name)
                        if os.path.isfile(fp):
                            expanded.append(fp)
                except OSError:
                    pass
            elif os.path.isfile(p):
                expanded.append(p)

        for fp in expanded:
            if fp in self.path_to_iid:
                continue
            try:
                size = os.path.getsize(fp)
            except OSError:
                size = 0
            iid = f"r{self._iid_counter}"
            self._iid_counter += 1
            try:
                kind = file_kind(fp)
            except Exception:
                kind = "Other"
            self.rows[iid] = {"path": fp, "size": size, "enabled": True}
            self.path_to_iid[fp] = iid
            self.tree.insert("", "end", iid=iid, tags=("ready",), image=self.img_on, text=" ON", values=(
                os.path.basename(fp), os.path.dirname(fp), kind,
                format_size(size), "Ready"))

        self._refresh_summary()

    def _active_iids(self):
        return [iid for iid, r in self.rows.items() if r["enabled"]]

    def _refresh_counts(self):
        """Update the summary line and the Start button (does not touch the progress bar)."""
        count = len(self.rows)
        active = [r for r in self.rows.values() if r["enabled"]]
        total = sum(r["size"] for r in active)
        self.summary_label.config(
            text=f"{count} file{'s' if count != 1 else ''}  ·  "
                 f"{len(active)} ON  ·  {format_size(total)} to convert")
        self.start_btn.set_enabled(bool(active) and not self.is_processing)

    def _apply_row_state(self, iid):
        info = self.rows[iid]
        self.results.pop(iid, None)
        if info["enabled"]:
            self.tree.item(iid, image=self.img_on, text=" ON")
            self._set_status(iid, "Ready", "ready")
        else:
            self.tree.item(iid, image=self.img_off, text=" OFF")
            self._set_status(iid, "Skipped", "off")

    def toggle_row(self, iid):
        if self.is_processing or iid not in self.rows:
            return
        self.rows[iid]["enabled"] = not self.rows[iid]["enabled"]
        self._apply_row_state(iid)
        self._refresh_counts()

    def set_all_enabled(self, enabled: bool):
        if self.is_processing:
            return
        for iid, info in self.rows.items():
            if info["enabled"] != enabled:
                info["enabled"] = enabled
                self._apply_row_state(iid)
        self._refresh_counts()

    def on_select_all(self):
        self.set_all_enabled(True)
        return "break"

    def on_deselect_all(self):
        self.set_all_enabled(False)

    def _on_tree_click(self, event):
        if self.is_processing:
            return None
        if self.tree.identify_region(event.x, event.y) not in ("tree", "cell"):
            return None
        if self.tree.identify_column(event.x) != "#0":
            return None
        iid = self.tree.identify_row(event.y)
        if iid in self.rows:
            self.toggle_row(iid)
            return "break"
        return None

    def _on_tree_space(self, _event):
        """Space bar: switch the selected rows ON/OFF together."""
        if self.is_processing:
            return "break"
        sel = [i for i in self.tree.selection() if i in self.rows]
        if sel:
            new_state = not all(self.rows[i]["enabled"] for i in sel)
            for iid in sel:
                if self.rows[iid]["enabled"] != new_state:
                    self.rows[iid]["enabled"] = new_state
                    self._apply_row_state(iid)
            self._refresh_counts()
        return "break"

    def _refresh_summary(self):
        count = len(self.rows)
        self._refresh_counts()
        if count:
            self.empty_label.place_forget()
        else:
            self.empty_label.place(relx=0.5, rely=0.5, anchor="center")
        self.progress["value"] = 0
        self.percent_label.config(text="0%")
        if not self.is_processing:
            self.status_label.config(text="Ready." if count else "No files selected.")

    def on_add_files(self):
        if self.is_processing:
            return
        paths = filedialog.askopenfilenames(
            title="Select files to convert",
            filetypes=[
                ("All files", "*.*"),
                ("Excel files", "*.xlsx *.xlsm *.xls"),
                ("CSV files", "*.csv"),
                ("Text/Log files", "*.txt *.log"),
                ("Database files", "*.db *.sqlite *.sqlite3 *.db3 *.mdb *.accdb *.mdf"),
                ("SQL dumps", "*.sql"),
                ("Documents", "*.pdf *.docx *.doc"),
                ("Torrent files", "*.torrent"),
            ])
        if paths:
            self.add_paths(paths)

    def on_add_folder(self):
        if self.is_processing:
            return
        folder = filedialog.askdirectory(title="Select a folder (all files inside will be added)")
        if folder:
            self.add_paths([folder])

    def on_remove_selected(self):
        if self.is_processing:
            return
        for iid in self.tree.selection():
            info = self.rows.pop(iid, None)
            if info:
                self.path_to_iid.pop(info["path"], None)
            self.results.pop(iid, None)
            self.tree.delete(iid)
        self._refresh_summary()

    def on_clear(self):
        if self.is_processing:
            return
        self.tree.delete(*self.tree.get_children())
        self.rows.clear()
        self.path_to_iid.clear()
        self.results.clear()
        self.open_btn.set_enabled(False)
        self.last_output = None
        self._refresh_summary()

    def _set_status(self, iid, text, tag):
        if iid in self.rows:
            self.tree.set(iid, "status", text)
            self.tree.item(iid, tags=(tag,))

    # ------------------------------------------------------------------
    # Running the conversion
    # ------------------------------------------------------------------
    def on_start(self):
        active = self._active_iids()
        if self.is_processing or not self.rows:
            return
        if not active:
            messagebox.showinfo(APP_NAME, "All files are switched OFF.\n\n"
                                "Switch at least one file ON (or press \u201cSelect All\u201d).")
            return
        out_dir = None
        if self.out_mode == "custom":
            if not self.out_dir:
                messagebox.showinfo(APP_NAME, "Choose an output folder first "
                                    "(\u201cA folder I choose\u2026\u201d) or switch back to "
                                    "\u201cNext to each file\u201d.")
                return
            try:
                os.makedirs(self.out_dir, exist_ok=True)
            except OSError as exc:
                messagebox.showerror("Output folder", f"Cannot use this folder:\n{self.out_dir}\n\n{exc}")
                return
            out_dir = self.out_dir
        self._run_out_dir = out_dir
        self.is_processing = True
        self.total_active = len(active)
        self.processed_files = 0
        self.success_count = 0
        self.partial.clear()
        self.results.clear()
        self.last_output = None
        for iid, info in self.rows.items():
            if info["enabled"]:
                self._set_status(iid, "Queued", "queued")
            else:
                self._set_status(iid, "Skipped", "off")
        self.progress["value"] = 0
        self.percent_label.config(text=f"0%  ·  0 / {self.total_active}")
        self.status_label.config(text="Starting…")
        for b in (self.start_btn, self.open_btn, self.clear_btn, self.remove_btn,
                  self.select_all_btn, self.deselect_all_btn, *self.out_btns):
            b.set_enabled(False)

        files = [self.rows[i]["path"] for i in active]
        threading.Thread(target=self._run_pool, args=(files, out_dir), daemon=True).start()

    def _run_pool(self, files, out_dir=None):
        total = len(files)

        def on_progress(path, pct):
            self.ui_queue.put({"type": "sql_progress", "path": path, "percent": pct})

        def on_done(path, success, out_path, error, done_count, total_count):
            self.ui_queue.put({"type": "file_done", "path": path, "success": success,
                               "out_path": out_path, "error": error, "total": total_count})

        try:
            run_batch_conversion(files, on_sql_progress=on_progress, on_file_done=on_done,
                                 out_dir=out_dir)
            self.ui_queue.put({"type": "all_done", "total": total})
        except Exception as exc:
            self.ui_queue.put({"type": "fatal", "error": f"{type(exc).__name__}: {exc}"})

    def _poll_queue(self):
        try:
            while True:
                self._handle_message(self.ui_queue.get_nowait())
        except queue.Empty:
            pass
        self.root.after(50, self._poll_queue)

    def _update_overall(self):
        total = self.total_active
        if not total:
            return
        value = (self.processed_files + sum(self.partial.values())) / total * 100
        self.progress["value"] = min(100, value)
        self.percent_label.config(text=f"{int(value)}%  ·  {self.processed_files} / {total}")

    def _handle_message(self, msg):
        kind = msg["type"]

        if kind == "drop":
            try:
                self.add_paths(msg["paths"])
            except Exception:
                _log_exception("add_paths (drop)")

        elif kind == "sql_progress":
            iid = self.path_to_iid.get(msg["path"])
            if iid:
                pct = float(msg["percent"])
                self.partial[iid] = max(0.0, min(1.0, pct / 100.0))
                self._set_status(iid, f"Parsing {pct:.0f}%", "working")
                self.status_label.config(text=f"Parsing {os.path.basename(msg['path'])}…")
                self._update_overall()

        elif kind == "file_done":
            iid = self.path_to_iid.get(msg["path"])
            self.processed_files += 1
            self.partial.pop(iid, None)
            if msg["success"]:
                self.success_count += 1
                if self.last_output is None:
                    self.last_output = msg["out_path"]
            if iid:
                self.results[iid] = (msg["success"], msg["out_path"], msg["error"])
                if msg["success"]:
                    self._set_status(iid, "✔  Done", "done")
                else:
                    self._set_status(iid, "✖  Failed", "failed")
            self.status_label.config(
                text=f"Processed {self.processed_files} of {msg['total']}…")
            self._update_overall()

        elif kind == "all_done":
            self._finish(msg["total"])

        elif kind == "fatal":
            self.is_processing = False
            for b in (self.clear_btn, self.remove_btn, self.select_all_btn, self.deselect_all_btn, *self.out_btns):
                b.set_enabled(True)
            self._refresh_counts()
            self.status_label.config(text="Conversion stopped because of an error.")
            messagebox.showerror("Error", msg["error"])

    def _finish(self, total):
        self.is_processing = False
        failed = total - self.success_count
        self.progress["value"] = 100
        self.percent_label.config(text=f"100%  ·  {total} / {total}")
        if failed == 0:
            self.status_label.config(text=f"All {total} file(s) converted successfully.")
        else:
            self.status_label.config(
                text=f"{self.success_count} succeeded, {failed} failed. Double-click a failed row for details.")
        for b in (self.clear_btn, self.remove_btn, self.select_all_btn, self.deselect_all_btn, *self.out_btns):
            b.set_enabled(True)
        self._refresh_counts()
        self.open_btn.set_enabled(self.last_output is not None)

        if failed == 0:
            messagebox.showinfo("Conversion completed",
                                f"All files were converted successfully.\n\nSucceeded: {self.success_count} / {total}"
                                f"\nSaved to: {self._output_where()}")
        else:
            lines = []
            for iid, (ok, _out, err) in self.results.items():
                if not ok:
                    lines.append(f"• {os.path.basename(self.rows[iid]['path'])}: {err}")
            shown = "\n".join(lines[:8])
            more = f"\n… and {len(lines) - 8} more" if len(lines) > 8 else ""
            messagebox.showwarning(
                "Conversion finished with errors",
                f"Succeeded: {self.success_count} / {total}    Failed: {failed}\n"
                f"Saved to: {self._output_where()}\n\n{shown}{more}")

    # ------------------------------------------------------------------
    # Table actions
    # ------------------------------------------------------------------
    def on_row_double_click(self, event):
        if self.tree.identify_column(event.x) == "#0":
            return  # the first column is the ON/OFF switch, not an "open" action
        iid = self.tree.identify_row(event.y)
        if not iid or iid not in self.rows:
            return
        result = self.results.get(iid)
        if result and not result[0]:
            messagebox.showerror(os.path.basename(self.rows[iid]["path"]),
                                 result[2] or "Unknown error")
        elif result and result[0] and result[1]:
            open_in_explorer(result[1])
        else:
            open_in_explorer(self.rows[iid]["path"])

    # ------------------------------------------------------------------
    # Output location
    # ------------------------------------------------------------------
    def _output_where(self) -> str:
        return self.out_dir if self.out_mode == "custom" else "next to each source file"

    def _refresh_output_ui(self):
        custom = self.out_mode == "custom"
        self.out_src_btn.kind = "ghost" if custom else "primary"
        self.out_dir_btn.kind = "primary" if custom else "ghost"
        for b in self.out_btns:
            b._paint()
        if custom:
            path = self.out_dir or "(no folder chosen yet)"
            if len(path) > 55:
                path = "\u2026" + path[-54:]
            self.out_label.config(text=path)
        else:
            self.out_label.config(text="Each CSV is saved in the same folder as its source file.")

    def _save_output_settings(self):
        self.settings["out_mode"] = self.out_mode
        self.settings["out_dir"] = self.out_dir
        save_settings(self.settings)

    def on_out_source(self):
        if self.is_processing:
            return
        self.out_mode = "source"
        self._save_output_settings()
        self._refresh_output_ui()

    def on_out_choose(self):
        if self.is_processing:
            return
        kw = {"title": "Choose the folder where the CSV files will be saved"}
        if self.out_dir and os.path.isdir(self.out_dir):
            kw["initialdir"] = self.out_dir
        folder = filedialog.askdirectory(**kw)
        if folder:
            self.out_dir = os.path.normpath(folder)
            self.out_mode = "custom"
            self._save_output_settings()
            self._refresh_output_ui()

    def on_open_output(self):
        if self.last_output:
            open_in_explorer(self.last_output)


def main():
    args = sys.argv[1:]
    out_dir = None
    if "--out" in args:  # --out FOLDER : save all CSVs in FOLDER instead of next to the sources
        i = args.index("--out")
        out_dir = args[i + 1] if i + 1 < len(args) else None
        args = args[:i] + args[i + 2:]
    cli = False
    if args and args[0] == "--cli":
        cli, args = True, args[1:]
    elif args and sys.stdout is not None and sys.stdout.isatty():
        cli = True  # run from a terminal with files: headless mode as before

    if cli:
        _run_cli_mode(args, out_dir)
        return

    # DPI awareness for a crisp display on high-resolution screens (Windows)
    if os.name == "nt":
        try:
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            try:
                ctypes.windll.user32.SetProcessDPIAware()
            except Exception:
                pass

    _enable_crash_log()
    root = tk.Tk()
    # an error inside any Tk callback must never close the app: log it and carry on
    root.report_callback_exception = lambda exc, val, tb: _log_exception("tk callback")
    BatchCsvConverterApp(root, initial_files=[a for a in args if os.path.exists(a)])
    root.mainloop()


if __name__ == "__main__":
    multiprocessing.freeze_support()  # required for the exe to work correctly on Windows (PyInstaller and similar)
    main()
